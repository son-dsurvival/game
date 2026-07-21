import json
import os
import secrets
import logging
from google.oauth2 import service_account
from googleapiclient.discovery import build
from googleapiclient.errors import HttpError
from io import BytesIO

from docx import Document
from fastapi import Depends, FastAPI, Header, HTTPException, Query
from pypdf import PdfReader


FOLDER_ID = os.getenv("GOOGLE_DRIVE_FOLDER_ID")
GOOGLE_SERVICE_ACCOUNT_JSON = os.getenv(
    "GOOGLE_SERVICE_ACCOUNT_JSON"
)
ACTION_API_KEY = os.getenv("ACTION_API_KEY")

SCOPES = [
    "https://www.googleapis.com/auth/drive.readonly"
]


if not FOLDER_ID:
    raise RuntimeError(
        "GOOGLE_DRIVE_FOLDER_ID environment variable is missing."
    )

if not GOOGLE_SERVICE_ACCOUNT_JSON:
    raise RuntimeError(
        "GOOGLE_SERVICE_ACCOUNT_JSON environment variable is missing."
    )

if not ACTION_API_KEY:
    raise RuntimeError(
        "ACTION_API_KEY environment variable is missing."
    )


try:
    service_account_info = json.loads(
        GOOGLE_SERVICE_ACCOUNT_JSON
    )
except json.JSONDecodeError as error:
    raise RuntimeError(
        "GOOGLE_SERVICE_ACCOUNT_JSON is not valid JSON."
    ) from error


credentials = service_account.Credentials.from_service_account_info(
    service_account_info,
    scopes=SCOPES,
)

drive_service = build(
    "drive",
    "v3",
    credentials=credentials,
    cache_discovery=False,
)


app = FastAPI(
    title="Google Drive Knowledge API",
    version="1.0.0",
)


def verify_api_key(
    authorization: str | None = Header(default=None),
) -> None:
    if not authorization:
        raise HTTPException(
            status_code=401,
            detail="Authorization header is missing.",
        )

    scheme, _, supplied_key = authorization.partition(" ")

    if (
        scheme.lower() != "bearer"
        or not supplied_key
        or not secrets.compare_digest(supplied_key, ACTION_API_KEY)
    ):
        raise HTTPException(
            status_code=401,
            detail="Invalid API key.",
        )
GOOGLE_DOC_MIME = "application/vnd.google-apps.document"
PDF_MIME = "application/pdf"
DOCX_MIME = (
    "application/vnd.openxmlformats-officedocument."
    "wordprocessingml.document"
)

TEXT_MIME_TYPES = {
    "text/plain",
    "text/markdown",
    "text/csv",
    "application/json",
}


def extract_pdf_text(file_bytes: bytes) -> str:
    reader = PdfReader(BytesIO(file_bytes))

    pages = [
        page.extract_text() or ""
        for page in reader.pages
    ]

    return "\n\n".join(pages).strip()


def extract_docx_text(file_bytes: bytes) -> str:
    document = Document(BytesIO(file_bytes))

    content: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            content.append(text)

    # Include text stored inside tables.
    for table in document.tables:
        for row in table.rows:
            cells = [
                cell.text.strip()
                for cell in row.cells
                if cell.text.strip()
            ]

            if cells:
                content.append(" | ".join(cells))

    return "\n".join(content).strip()


def extract_text_file(file_bytes: bytes) -> str:
    return file_bytes.decode(
        "utf-8-sig",
        errors="replace",
    ).strip()


def get_drive_file_text(file_id: str) -> tuple[dict, str]:
    metadata = (
        drive_service.files()
        .get(
            fileId=file_id,
            fields=(
                "id,"
                "name,"
                "mimeType,"
                "parents,"
                "webViewLink,"
                "modifiedTime"
            ),
        )
        .execute()
    )

    # Prevent arbitrary file IDs from being used to read other files
    # available to the service account.
    parents = metadata.get("parents", [])

    if FOLDER_ID not in parents:
        raise HTTPException(
            status_code=403,
            detail=(
                "The requested file is not directly inside "
                "the authorised Google Drive folder."
            ),
        )

    mime_type = metadata.get("mimeType", "")

    if mime_type == GOOGLE_DOC_MIME:
        file_bytes = (
            drive_service.files()
            .export_media(
                fileId=file_id,
                mimeType="text/plain",
            )
            .execute()
        )

        text = extract_text_file(file_bytes)

    elif mime_type == PDF_MIME:
        file_bytes = (
            drive_service.files()
            .get_media(fileId=file_id)
            .execute()
        )

        text = extract_pdf_text(file_bytes)

    elif mime_type == DOCX_MIME:
        file_bytes = (
            drive_service.files()
            .get_media(fileId=file_id)
            .execute()
        )

        text = extract_docx_text(file_bytes)

    elif mime_type in TEXT_MIME_TYPES:
        file_bytes = (
            drive_service.files()
            .get_media(fileId=file_id)
            .execute()
        )

        text = extract_text_file(file_bytes)

    else:
        raise HTTPException(
            status_code=415,
            detail=f"Unsupported file type: {mime_type}",
        )

    if not text:
        raise HTTPException(
            status_code=422,
            detail=(
                "The file was retrieved, but no readable text "
                "could be extracted."
            ),
        )

    return metadata, text
logger = logging.getLogger("uvicorn.error")
@app.get("/")
def home():
    return {
        "status": "running",
        "message": "Google Drive Knowledge API is running.",
    }


@app.get("/files")
def list_drive_files(
    _: None = Depends(verify_api_key),
):
    try:
        query = f"'{FOLDER_ID}' in parents and trashed = false"

        response = (
            drive_service.files()
            .list(
                q=query,
                spaces="drive",
                pageSize=100,
                fields=(
                    "files("
                    "id,"
                    "name,"
                    "mimeType,"
                    "modifiedTime,"
                    "webViewLink"
                    ")"
                ),
            )
            .execute()
        )

        files = response.get("files", [])

        return {
            "folder_id": FOLDER_ID,
            "count": len(files),
            "files": files,
        }

    except HttpError as error:
        raise HTTPException(
            status_code=502,
            detail=f"Google Drive API error: {error}",
        ) from error
GOOGLE_DOC_MIME = "application/vnd.google-apps.document"


@app.get("/documents/{file_id}")
def read_drive_document(
    file_id: str,
    start: int = Query(default=0, ge=0),
    max_chars: int = Query(
        default=12000,
        ge=1000,
        le=30000,
    ),
    _: None = Depends(verify_api_key),
):
    try:
        metadata = (
            drive_service.files()
            .get(
                fileId=file_id,
                fields=(
                    "id,"
                    "name,"
                    "mimeType,"
                    "parents,"
                    "webViewLink,"
                    "modifiedTime"
                ),
            )
            .execute()
        )

        parents = metadata.get("parents", [])

        if FOLDER_ID not in parents:
            raise HTTPException(
                status_code=403,
                detail="File is outside the authorised folder.",
            )

        mime_type = metadata.get("mimeType")

        if mime_type != GOOGLE_DOC_MIME:
            raise HTTPException(
                status_code=415,
                detail=f"Unsupported file type: {mime_type}",
            )

        exported_bytes = (
            drive_service.files()
            .export_media(
                fileId=file_id,
                mimeType="text/plain",
            )
            .execute()
        )

        if not isinstance(exported_bytes, (bytes, bytearray)):
            raise RuntimeError(
                f"Unexpected export result: {type(exported_bytes).__name__}"
            )

        text = exported_bytes.decode(
            "utf-8",
            errors="replace",
        ).strip()

        if not text:
            raise HTTPException(
                status_code=422,
                detail="The document contains no readable text.",
            )

        if start >= len(text):
            raise HTTPException(
                status_code=416,
                detail="Start position exceeds document length.",
            )

        end = min(start + max_chars, len(text))

        return {
            "file_id": metadata["id"],
            "title": metadata["name"],
            "mime_type": metadata["mimeType"],
            "modified_time": metadata.get("modifiedTime"),
            "url": metadata.get("webViewLink"),
            "content": text[start:end],
            "start": start,
            "end": end,
            "total_characters": len(text),
            "truncated": end < len(text),
            "next_start": end if end < len(text) else None,
        }

    except HTTPException:
        raise

    except HttpError as error:
        logger.exception("Google Drive export failed")

        raise HTTPException(
            status_code=502,
            detail={
                "error_type": "GoogleDriveHttpError",
                "google_status": error.resp.status,
                "message": str(error),
            },
        ) from error

    except Exception as error:
        logger.exception("Document reading failed")

        raise HTTPException(
            status_code=500,
            detail={
                "error_type": type(error).__name__,
                "message": str(error),
            },
        ) from error
